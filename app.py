from flask import (Flask, render_template, request, redirect,
                   url_for, session, send_file, jsonify, abort) 
from groq import Groq
import os, datetime, io, uuid, json

# ── Optional deps ─────────────────────────────────────────────────────────────
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "helpdesk-secret-2024")

# ── Config ────────────────────────────────────────────────────────────────────
ADMIN_USERNAME  = os.environ.get("ADMIN_USERNAME", "admin")
ADMIN_PASSWORD  = os.environ.get("ADMIN_PASSWORD", "admin123")
UPLOAD_FOLDER   = os.path.join(os.path.dirname(__file__), "static", "uploads")
ALLOWED_EXT     = {"png", "jpg", "jpeg", "gif", "pdf", "txt", "docx", "xlsx"}
SLA_HOURS       = int(os.environ.get("SLA_HOURS", 24))   # breach after N hours

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

client = Groq(api_key=os.environ.get("GROQ_API_KEY"))

# ── In-memory store ───────────────────────────────────────────────────────────
tickets = []
ticket_counter = 1


# ── Helpers ───────────────────────────────────────────────────────────────────
def login_required(f):
    from functools import wraps
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXT


def auto_priority(issue):
    low = issue.lower()
    if any(k in low for k in ["urgent","crash","down","broken","emergency","critical","not working","error"]):
        return "High"
    if any(k in low for k in ["slow","minor","question","how to","update","suggestion"]):
        return "Low"
    return "Medium"


def get_ai_suggestion(issue, category):
    try:
        prompt = (
            f"You are a helpful IT support assistant. A user submitted a helpdesk ticket.\n"
            f"Category: {category}\nIssue: {issue}\n\n"
            f"Give a concise, friendly troubleshooting suggestion in 2-3 sentences. Be specific."
        )
        resp = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[{"role": "user", "content": prompt}],
            max_tokens=150,
        )
        return resp.choices[0].message.content.strip()
    except Exception as e:
        return f"AI suggestion unavailable: {e}"


def sla_status(ticket):
    """Return dict with hours_elapsed, breached, percent (0-100), time_left_str."""
    try:
        created = datetime.datetime.strptime(ticket["created_at"], "%Y-%m-%d %H:%M")
    except Exception:
        return {"hours_elapsed": 0, "breached": False, "percent": 0, "time_left_str": "N/A"}

    now     = datetime.datetime.now()
    elapsed = (now - created).total_seconds() / 3600          # in hours
    breached = elapsed >= SLA_HOURS
    percent  = min(100, int((elapsed / SLA_HOURS) * 100))

    if not breached:
        left_h = int(SLA_HOURS - elapsed)
        left_m = int(((SLA_HOURS - elapsed) - left_h) * 60)
        time_left_str = f"{left_h}h {left_m}m left"
    else:
        over_h = int(elapsed - SLA_HOURS)
        time_left_str = f"Breached by {over_h}h"

    return {
        "hours_elapsed": round(elapsed, 1),
        "breached": breached,
        "percent": percent,
        "time_left_str": time_left_str,
    }


def build_chart_data():
    """Return JSON-serialisable dicts for the dashboard charts."""
    status_counts   = {}
    priority_counts = {}
    category_counts = {}
    # tickets by day (last 7 days)
    day_counts = {}
    for i in range(6, -1, -1):
        d = (datetime.datetime.now() - datetime.timedelta(days=i)).strftime("%b %d")
        day_counts[d] = 0

    for t in tickets:
        status_counts[t["status"]]     = status_counts.get(t["status"], 0) + 1
        priority_counts[t["priority"]] = priority_counts.get(t["priority"], 0) + 1
        cat = t.get("category", "General")
        category_counts[cat] = category_counts.get(cat, 0) + 1
        # trend
        try:
            d = datetime.datetime.strptime(t["created_at"], "%Y-%m-%d %H:%M").strftime("%b %d")
            if d in day_counts:
                day_counts[d] += 1
        except Exception:
            pass

    return {
        "status":   {"labels": list(status_counts.keys()),   "data": list(status_counts.values())},
        "priority": {"labels": list(priority_counts.keys()), "data": list(priority_counts.values())},
        "category": {"labels": list(category_counts.keys()), "data": list(category_counts.values())},
        "trend":    {"labels": list(day_counts.keys()),      "data": list(day_counts.values())},
    }


# ── Auth ──────────────────────────────────────────────────────────────────────
@app.route("/login", methods=["GET","POST"])
def login():
    error = None
    if request.method == "POST":
        u = request.form.get("username","").strip()
        p = request.form.get("password","")
        if u == ADMIN_USERNAME and p == ADMIN_PASSWORD:
            session["logged_in"] = True
            session["username"]  = u
            return redirect(url_for("admin"))
        error = "Invalid username or password."
    return render_template("login.html", error=error)

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ── User: submit ticket ───────────────────────────────────────────────────────
@app.route("/")
def index():
    return render_template("index.html")

@app.route("/submit", methods=["POST"])
def submit_ticket():
    global ticket_counter
    name     = request.form.get("name","").strip()
    email    = request.form.get("email","").strip()
    issue    = request.form.get("issue","").strip()
    category = request.form.get("category","General")

    if not name or not issue:
        return render_template("index.html", error="Name and issue are required.")

    # ── File attachment ───────────────────────────────────────────────────────
    attachment = None
    file = request.files.get("attachment")
    if file and file.filename and allowed_file(file.filename):
        ext      = file.filename.rsplit(".", 1)[1].lower()
        filename = f"{uuid.uuid4().hex}.{ext}"
        file.save(os.path.join(UPLOAD_FOLDER, filename))
        attachment = {"original": file.filename, "saved": filename, "ext": ext}

    ai_suggestion = get_ai_suggestion(issue, category)

    now = datetime.datetime.now()
    ticket = {
        "id":           ticket_counter,
        "name":         name,
        "email":        email,
        "issue":        issue,
        "category":     category,
        "ai_suggestion":ai_suggestion,
        "status":       "Open",
        "priority":     auto_priority(issue),
        "created_at":   now.strftime("%Y-%m-%d %H:%M"),
        "created_date": now.strftime("%Y-%m-%d"),
        "attachment":   attachment,
        "comments":     [],          # list of {author, text, timestamp}
        "assigned_to":  None,
        "sla_breached": False,
    }
    tickets.append(ticket)
    ticket_counter += 1

    return render_template("index.html", ticket=ticket)


# ── Admin: main panel ─────────────────────────────────────────────────────────
@app.route("/admin")
@login_required
def admin():
    status_filter   = request.args.get("status",   "All")
    priority_filter = request.args.get("priority", "All")
    search          = request.args.get("search",   "").strip().lower()

    # Update SLA breach flags
    for t in tickets:
        if t["status"] not in ("Resolved","Closed"):
            s = sla_status(t)
            t["sla_breached"] = s["breached"]
            if s["breached"] and t["priority"] != "High":
                t["priority"] = "High"   # auto-escalate

    filtered = tickets[:]
    if status_filter   != "All": filtered = [t for t in filtered if t["status"]   == status_filter]
    if priority_filter != "All": filtered = [t for t in filtered if t["priority"] == priority_filter]
    if search:
        filtered = [t for t in filtered if
                    search in t["name"].lower() or
                    search in t["issue"].lower() or
                    search in t.get("email","").lower()]

    stats = {
        "total":    len(tickets),
        "open":     sum(1 for t in tickets if t["status"] == "Open"),
        "resolved": sum(1 for t in tickets if t["status"] == "Resolved"),
        "high":     sum(1 for t in tickets if t["priority"] == "High"),
        "breached": sum(1 for t in tickets if t.get("sla_breached")),
    }

    # Attach sla info to each ticket for the template
    for t in filtered:
        t["_sla"] = sla_status(t)

    chart_data = json.dumps(build_chart_data())

    return render_template("admin.html",
                           tickets=filtered, stats=stats,
                           status_filter=status_filter,
                           priority_filter=priority_filter,
                           search=search,
                           excel_available=EXCEL_AVAILABLE,
                           chart_data=chart_data,
                           sla_hours=SLA_HOURS)


# ── Admin: ticket detail ──────────────────────────────────────────────────────
@app.route("/admin/ticket/<int:ticket_id>")
@login_required
def ticket_detail(ticket_id):
    t = next((x for x in tickets if x["id"] == ticket_id), None)
    if not t:
        abort(404)
    t["_sla"] = sla_status(t)
    return render_template("ticket_detail.html", ticket=t, sla_hours=SLA_HOURS)


# ── Admin: add comment ────────────────────────────────────────────────────────
@app.route("/admin/ticket/<int:ticket_id>/comment", methods=["POST"])
@login_required
def add_comment(ticket_id):
    t = next((x for x in tickets if x["id"] == ticket_id), None)
    if not t:
        abort(404)
    text = request.form.get("comment","").strip()
    if text:
        t["comments"].append({
            "author":    session.get("username","Admin"),
            "text":      text,
            "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
            "type":      request.form.get("comment_type","internal"),  # internal | reply
        })
    return redirect(url_for("ticket_detail", ticket_id=ticket_id))


# ── Admin: update status ──────────────────────────────────────────────────────
@app.route("/admin/update/<int:ticket_id>", methods=["POST"])
@login_required
def update_ticket(ticket_id):
    t = next((x for x in tickets if x["id"] == ticket_id), None)
    if t:
        t["status"]      = request.form.get("status", t["status"])
        t["assigned_to"] = request.form.get("assigned_to", t.get("assigned_to","")).strip() or None
    ref = request.form.get("ref","admin")
    return redirect(url_for("ticket_detail", ticket_id=ticket_id) if ref=="detail" else url_for("admin"))


# ── Admin: delete ─────────────────────────────────────────────────────────────
@app.route("/admin/delete/<int:ticket_id>", methods=["POST"])
@login_required
def delete_ticket(ticket_id):
    global tickets
    tickets = [t for t in tickets if t["id"] != ticket_id]
    return redirect(url_for("admin"))


# ── Admin: dashboard charts API ───────────────────────────────────────────────
@app.route("/admin/chart-data")
@login_required
def chart_data_api():
    return jsonify(build_chart_data())


# ── File download ─────────────────────────────────────────────────────────────
@app.route("/uploads/<filename>")
@login_required
def uploaded_file(filename):
    path = os.path.join(UPLOAD_FOLDER, filename)
    if not os.path.exists(path):
        abort(404)
    return send_file(path)


# ── Excel export ──────────────────────────────────────────────────────────────
@app.route("/admin/export/excel")
@login_required
def export_excel():
    if not EXCEL_AVAILABLE:
        return redirect(url_for("admin"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HelpDesk Tickets"

    header_font  = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    header_fill  = PatternFill("solid", fgColor="4C3FB5")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    thin         = Side(border_style="thin", color="DDDDDD")
    cell_border  = Border(left=thin, right=thin, top=thin, bottom=thin)
    priority_colors = {"High":"FFCDD2","Medium":"FFF9C4","Low":"C8E6C9"}
    status_colors   = {"Open":"BBDEFB","In Progress":"FFE0B2","Resolved":"C8E6C9","Closed":"E0E0E0"}

    ws.merge_cells("A1:J1")
    ws["A1"].value     = f"HelpDesk Ticket Report — {datetime.datetime.now().strftime('%d %B %Y, %H:%M')}"
    ws["A1"].font      = Font(name="Calibri", bold=True, size=14, color="2C1B6B")
    ws["A1"].fill      = PatternFill("solid", fgColor="EDE9FF")
    ws["A1"].alignment = center_align
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:J2")
    ws["A2"].value     = (f"Total: {len(tickets)}  |  Open: {sum(1 for t in tickets if t['status']=='Open')}"
                          f"  |  Resolved: {sum(1 for t in tickets if t['status']=='Resolved')}"
                          f"  |  SLA Breached: {sum(1 for t in tickets if t.get('sla_breached'))}")
    ws["A2"].font      = Font(name="Calibri", italic=True, size=10, color="555555")
    ws["A2"].alignment = center_align
    ws.row_dimensions[2].height = 18

    headers    = ["ID","Date","Name","Email","Category","Priority","Status","SLA","Assigned To","Issue"]
    col_widths = [6,  14,    18,    24,     14,        10,        12,      14,   16,            50]
    for ci,(h,w) in enumerate(zip(headers,col_widths),1):
        c = ws.cell(row=3, column=ci, value=h)
        c.font=header_font; c.fill=header_fill; c.alignment=center_align; c.border=cell_border
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[3].height = 22

    for ri,ticket in enumerate(tickets, 4):
        sla = sla_status(ticket)
        row = [ticket["id"], ticket.get("created_date",""), ticket["name"],
               ticket.get("email",""), ticket.get("category",""), ticket["priority"],
               ticket["status"], "BREACHED" if sla["breached"] else sla["time_left_str"],
               ticket.get("assigned_to","—"), ticket["issue"]]
        for ci,val in enumerate(row,1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.border=cell_border; c.font=Font(name="Calibri",size=10)
            c.alignment = left_align if ci==10 else center_align
            if ci==6:
                c.fill = PatternFill("solid", fgColor=priority_colors.get(ticket["priority"],"FFFFFF"))
            if ci==7:
                c.fill = PatternFill("solid", fgColor=status_colors.get(ticket["status"],"FFFFFF"))
            if ci==8 and sla["breached"]:
                c.fill = PatternFill("solid", fgColor="FFCDD2")
                c.font = Font(name="Calibri", size=10, bold=True, color="C62828")
        ws.row_dimensions[ri].height = 20

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:J{3+len(tickets)}"

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    fname = f"helpdesk_tickets_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    app.run(debug=True, port=5000)
